import pandas as pd, numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
E=pd.read_pickle('events.pkl').sort_values('d5')
pct=lambda v: None if v is None or v!=v else round(float(v)*100,1)
yn=lambda v: '是' if v is True or v==1.0 else ''
H=['股票','代码','五板日',
   '第二天','开盘%','最高%','最低%','收盘%','碰涨停','碰跌停','收涨停','收跌停','开盘买到当日收%','开盘买到后天收%','收盘买到次日收%',
   '第三天','开盘%','最高%','最低%','收盘%','碰涨停','碰跌停','收涨停','收跌停','开盘买到当日收%','开盘买到后天收%','收盘买到次日收%']
wb=Workbook(); ws=wb.active; ws.title='五板后两日'
ws.append(['2026年7月以来五板股, 五板后第二天与第三天的开盘、盘中、收盘(相对前一日收盘), 以及在当天开盘或收盘买入后的收益; 不含10月19日盘中'])
ws.append([]); ws.append(H)
for _,r in E.iterrows():
    row=[r['name'],r.code,r.d5]
    for d in ('d2','d3'):
        if r.get(f'{d}_date')!=r.get(f'{d}_date'): row+=['']*12; continue
        row+=[r[f'{d}_date'],pct(r[f'{d}_gap']),pct(r[f'{d}_hi']),pct(r[f'{d}_lo']),pct(r[f'{d}_chg']),yn(r[f'{d}_touch_zt']),yn(r[f'{d}_touch_dt']),yn(r[f'{d}_close_zt']),yn(r[f'{d}_close_dt']),
              pct(r[f'{d}_开盘买_当日收']),pct(r[f'{d}_开盘买_第3日收']),pct(r[f'{d}_收盘买_次日收'])]
    ws.append(row)
ws.append([]); ws.append(['来源: 新浪不复权日K(腾讯日K当日被拦); 五板=价格口径纯连续第5个涨停, 10cm/20cm按板块, ST自7月6日起按10%; 一字涨停开盘买不到记空; 后天收=买入日之后第2个交易日收盘'])
f=Font(name='Times New Roman',size=11); fb=Font(name='Times New Roman',size=11,bold=True); al=Alignment(horizontal='center',vertical='center')
for row in ws.iter_rows():
    for c in row: c.font=f; c.alignment=al
for c in ws[3]: c.font=fb
ws['A1'].alignment=Alignment(horizontal='left',vertical='center'); ws.cell(ws.max_row,1).alignment=Alignment(horizontal='left',vertical='center')
ws.sheet_view.showGridLines=False; ws.freeze_panes='D4'
for i,w in enumerate([10,8,11]+[11,7,7,7,7,6,6,6,6,10,10,10]*2,1): ws.column_dimensions[ws.cell(3,i).column_letter].width=w
wb.save('../../output/五板后两日.xlsx'); print('ok', len(E))
