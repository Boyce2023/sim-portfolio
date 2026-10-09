#!/opt/homebrew/bin/python3.13
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
X=pd.read_csv('signals_features.csv',parse_dates=['date']).sort_values('date')
F=Font(name='Times New Roman',size=11); B=Font(name='Times New Roman',size=11,bold=True)
wb=Workbook()
def sheet(ws,rows,widths):
    ws.sheet_view.showGridLines=False; ws.column_dimensions['A'].width=1.83
    for r,row in enumerate(rows,start=2):
        for c,v in enumerate(row,start=2):
            cell=ws.cell(r,c,v); cell.font=F; cell.alignment=Alignment(vertical='center')
    from openpyxl.utils import get_column_letter
    for i,w in enumerate(widths,start=2): ws.column_dimensions[get_column_letter(i)].width=w
def bold_row(ws,r):
    for c in ws[r]: 
        if c.value is not None: c.font=B
ws=wb.active; ws.title='结论'
def stat(m,label):
    x=X[m]; a=x[x.段=='前段']; b=x[x.段=='后段']; e=x.ex10.dropna()
    return [label,len(x),x.命中.mean(),a.命中.mean(),b.命中.mean(),(x.rank10==1).mean(),(e>0).mean(),e.mean(),e.median()]
C=[(X.index==X.index,'全部信号'),(X.当日排名<=2,'A 信号当日该题材涨停数排全市场前二'),(X.二板以上只数>=4,'B 该题材当日二板以上>=4只'),
   (X.全市场涨停<=63,'C 全市场当日涨停<=63只(大盘弱时题材仍强)'),(X.前10日已进前三天数>=5,'D 前10日已有>=5天进前三(延续中的主线)'),
   ((X.当日排名<=2)&(X.二板以上只数>=4),'A+B'),((X.当日排名<=2)&(X.全市场涨停<=63),'A+C'),((X.当日排名<=2)&(X.前10日已进前三天数>=5),'A+D'),
   ((X.当日排名<=2)&(X.二板以上只数>=4)&(X.前10日已进前三天数>=5),'A+B+D')]
rows=[['题材连续两天强势：之后10日排进前二的概率与筛选条件'],[],
      ['结论：筛选能把"10日内排进前二"的命中率从61%提到76%到80%，前后两段都稳定；但次日开盘买入该题材当日涨停股，即使题材命中，10日平均仍跑输大盘，题材对不等于买入那几只票赚钱。'],[],
      ['筛选条件','信号数','命中率(10日内前二)','前段命中率(2025-10到2026-03)','后段命中率(2026-04到2026-09)','10日内第一','买入10日超额胜率','买入10日超额均值','买入10日超额中位']]
rows+=[stat(m,l) for m,l in C]
rows+=[[],['口径与覆盖'],
 ['数据','同花顺每日涨停池的涨停原因(当日标注，无事后归类)，2025-09-24到2026-10-09，接口只保留约一年'],
 ['题材','涨停原因按+拆词，同义词合并(固态电池/农业/医药/AI制药等)，剔除国资、业绩、回购等非题材词'],
 ['信号','某题材前一日与当日涨停只数都排全市场前三、每日至少5只、当日有二板以上'],
 ['命中','信号日之后10个交易日累计涨停只数排全市场前二'],
 ['买入','信号次日开盘买入该题材当日涨停股(开盘一字买不进的剔除)，持有10个交易日收盘，减同期全市场等权；价格库到2026-09-30，之后的信号无收益'],
 ['局限','样本约一年、191个信号；题材合并口径会影响排名；筛选阈值按三分位切，未做第三段样本外检验']]
sheet(ws,rows,[44,9,18,26,26,11,16,16,16])
bold_row(ws,2); bold_row(ws,6); bold_row(ws,6+len(C)+2)
for r in range(7,7+len(C)):
    for c in range(4,11): ws.cell(r,c).number_format='0%' if c<=8 else '0.0%'
    for c in (9,10): ws.cell(r,c).number_format='0.0%'
ws2=wb.create_sheet('信号明细')
cols=['date','theme','c_t1','c_t','当日排名','最高板','二板以上只数','炸板率','早封占比','一字占比','流通市值中位','换手中位','全市场涨停','前10日已进前三天数','rank5','rank10','zt10','命中','ex5','ex10','段']
names=['信号日','题材','前一日涨停数','当日涨停数','当日排名','最高板','二板以上只数','炸板率','早封占比','一字占比','涨停股流通市值中位(亿)','换手率中位(%)','全市场涨停数','前10日进前三天数','5日累计排名','10日累计排名','10日累计涨停数','命中(10日前二)','买入5日超额','买入10日超额','样本段']
d=X[cols].copy(); d['date']=d.date.dt.strftime('%Y-%m-%d')
sheet(ws2,[names]+d.values.tolist(),[12,14]+[12]*19); bold_row(ws2,2)
for r in range(3,len(d)+3):
    for c,fmt in [(9,'0%'),(10,'0%'),(11,'0%'),(12,'0.0'),(13,'0.0'),(20,'0.0%'),(21,'0.0%')]: ws2.cell(r,c).number_format=fmt
ws2.freeze_panes='D3'
wb.save('../../output/题材强势回测.xlsx'); print('ok')
