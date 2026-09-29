#!/opt/homebrew/bin/python3.13
"""交付: output/通达型前推回测.xlsx  (格式按 ~/.claude/standards/buwen_xlsx_format.md)"""
import pandas as pd, numpy as np, json, os, shutil, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
OUT='/Users/huaichuaibeimeng/claude-projects/sim-portfolio/output/通达型前推回测.xlsx'
V={'E':'C 重训+个股记忆(最终版)','A':'A 不学习','B':'B 每月重训','C':'C 加自校准闸门','D':'C 加闸门+卖出参数月调','R':'R 同池随机'}
names=pd.read_csv('names.csv',dtype={'code':str}); names['code']=names.code.str.zfill(6); nm=dict(zip(names.code,names.name))
POOL={1:'横盘点火',2:'台阶初期'}
F=pd.read_parquet('feat.parquet',columns=['date','code','pool'])
pool=F.set_index(['date','code']).pool
def trades(k):
    T=pd.read_csv(f'trades_{k}.csv',parse_dates=['buy_d','sell_d'],dtype={'code':str})
    T['code']=T.code.str.zfill(6)
    return T
wb=Workbook(); f11=Font(name='Times New Roman',size=11); fb=Font(name='Times New Roman',size=11,bold=True)
def sheet(title,blocks,widths=None):
    ws=wb.create_sheet(title); ws.sheet_view.showGridLines=False; ws.column_dimensions['A'].width=1.83
    r=2
    for head,df in blocks:
        if head: ws.cell(r,2,head).font=fb; r+=1
        if isinstance(df,str):
            for line in df.split('\n'): ws.cell(r,2,line).font=f11; r+=1
            r+=1; continue
        for j,c in enumerate(df.columns): ws.cell(r,2+j,str(c)).font=fb
        r+=1
        for row in df.itertuples(index=False):
            for j,v in enumerate(row):
                if isinstance(v,(float,np.floating)) and np.isnan(v): v=None
                if isinstance(v,(pd.Timestamp,)): v=v.to_pydatetime().date()
                c=ws.cell(r,2+j,v.item() if hasattr(v,'item') else v); c.font=f11
            r+=1
        r+=1
    for col in ws.iter_cols(min_col=2,max_col=ws.max_column):
        L=max((len(str(c.value)) if c.value is not None else 0) for c in col)
        ws.column_dimensions[col[0].column_letter].width=min(max(8,L*1.6+2),60)
    for row in ws.iter_rows():
        for c in row: c.alignment=Alignment(vertical='center',horizontal=c.alignment.horizontal,wrap_text=False)
    return ws
S=pd.read_csv('summary.csv'); Q=pd.read_csv('by_quarter.csv',header=[0,1],index_col=0)
half=pd.read_csv('halves.csv'); wf=pd.DataFrame(json.load(open('wf_log.json')))
concl=open('conclusion.txt').read()
rp=pd.read_csv('robust_picks.csv'); rr=pd.read_csv('robust_runs.csv')
rq=rr[rr['类']=='R'].净值; rsum=pd.DataFrame([{'随机次数':len(rq),'中位净值':round(rq.median(),3),'5%分位':round(rq.quantile(.05),3),'95%分位':round(rq.quantile(.95),3),'最高':round(rq.max(),3)}])
r2=pd.read_csv('robust2.csv'); pert=r2[['配置','N','K','净值','回撤','笔数','胜率','每笔期望','赔率','超过随机%']].rename(columns={'N':'仓位数','K':'每日新开'}).round(3); pert['配置']=pert['配置'].map({'sA无闸门':'A 不学习','sB无闸门':'B 每月重训','sC无闸门':'C 重训+个股记忆'})
sheet('结论',[('结论',concl),('各版本全期表现(2024-10-08 至 2026-09-23)',S),('前后两半段',half),('同池随机对照分布',rsum),('三个模型 × 9 种仓位组合(均不加闸门)',pert),('逐日选股质量(不经组合路径, 次日开盘买持有10日, 已扣费)',rp)])
q=Q.copy(); q.columns=[f'{a} {b}' for a,b in q.columns]; q.insert(0,'季度',q.index.astype(str))
sheet('分季度',[('按买入季度统计: 胜率%/每笔期望%/赔率',q.round(2))])
wf2=wf.rename(columns={'month':'训练日','train_rows':'训练样本','rows':'当月打分样本','sA':'IC A','sB':'IC B','sC':'IC C','sec':'耗时秒'}).round(3)
sheet('逐月学习记录',[('每月初重训一次, 只用标签已揭晓的样本; IC=当月分数与之后10日实际超额的秩相关(事后评估, 不回灌模型)',wf2)])
def tt(k):
    T=trades(k).copy()
    T['名称']=T.code.map(nm)
    sig=[]
    for _,t in T.iterrows(): sig.append(None)
    T=T.rename(columns={'code':'代码','buy_d':'买入日','sell_d':'卖出日','ret':'收益%','days':'持有天数','mfe':'最大浮盈%','mae':'最大浮亏%','why':'卖出原因',
        'rank':'当日排名','r5':'5日涨幅%','r20':'20日涨幅%','box20':'20日箱体%','t20':'20日均换手%','zt20':'20日涨停数','fcap':'流通市值亿','score':'模型分数'})
    for c in ['收益%','最大浮盈%','最大浮亏%','5日涨幅%','20日涨幅%','20日箱体%','20日均换手%']: T[c]=(T[c]*100).round(2)
    T['流通市值亿']=T['流通市值亿'].round(1); T['模型分数']=T['模型分数'].round(4)
    T['结果']=np.where(T['收益%']>0,'对','错')
    T['通达型']=np.where(T['最大浮盈%']>=25,'是','')
    return T[['代码','名称','买入日','卖出日','持有天数','收益%','结果','最大浮盈%','最大浮亏%','通达型','卖出原因','当日排名','模型分数','5日涨幅%','20日涨幅%','20日箱体%','20日均换手%','20日涨停数','流通市值亿']]
sheet('交易_最终版',[('最终版全部交易: 信号日收盘选股, 次日开盘买; 卖出条件收盘触发, 再下一日开盘卖; 已含费用',tt('E'))])
for k in ['A','B','C','D','R']:
    if os.path.exists(f'trades_{k}.csv'): sheet(f'交易_{k}',[(V[k]+' 全部交易',tt(k))])
if os.path.exists('d_schedule.json'):
    d=json.load(open('d_schedule.json')); dd=pd.DataFrame([(m,*v) for m,v in d.items()],columns=['月份','止损','持有天数上限','浮盈启动线','回撤卖出'])
    sheet('D参数月调',[('每月初用过去126个交易日影子回测18组参数选最优(只用当时已发生的行情)',dd)])
wb.remove(wb['Sheet'])
if os.path.exists(OUT):
    arc=os.path.join(os.path.dirname(OUT),'_归档_底稿与旧版'); os.makedirs(arc,exist_ok=True)
    shutil.move(OUT,os.path.join(arc,f'通达型前推回测_{dt.datetime.now():%Y%m%d%H%M}.xlsx'))
wb.save(OUT); print('写出',OUT)
