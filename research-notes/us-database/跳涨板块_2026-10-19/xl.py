"""Buwen 选股活表格式: Times New Roman 11, 表头只加粗, 无填充无边框, 网格线关, A列1.83留白, 百分比存小数0.0%, 收益列0为中点红白绿色阶, 最右备注列灰字8A857C"""
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter as L
FN='Times New Roman'; GRAY='8A857C'; YEL=PatternFill('solid',fgColor='FFFF00')
FMT={'pct':'0.0%','int':'#,##0','px':'#,##0.00','pe':'0.0','x':'0.0','text':'General'}
def put(ws,r,c,v,k='text',bold=False,gray=False,fill=None):
    if isinstance(v,float) and v!=v: v=None
    x=ws.cell(r,c,v); x.font=Font(name=FN,size=11,bold=bold,color=GRAY if gray else None)
    x.number_format=FMT[k]; x.alignment=Alignment(vertical='center',wrap_text=False)
    if fill: x.fill=fill
    return x
def sheet(wb,name,title,caliber,cols,widths):
    ws=wb.create_sheet(name); ws.sheet_view.showGridLines=False; ws.column_dimensions['A'].width=1.83
    put(ws,1,2,title,bold=True); put(ws,2,2,caliber)
    for j,(h,_) in enumerate(cols): put(ws,4,2+j,h,bold=True)
    for j,w in enumerate(widths): ws.column_dimensions[L(2+j)].width=w
    ws.freeze_panes='A5'
    return ws
def scale(ws,col,r0,r1):
    ws.conditional_formatting.add(f'{L(col)}{r0}:{L(col)}{r1}',ColorScaleRule(start_type='min',start_color='F8696B',mid_type='num',mid_value=0,mid_color='FFFFFF',end_type='max',end_color='63BE7B'))
def notes(ws,r,lines):
    for k,s in enumerate(lines): put(ws,r+k,2,s,gray=True)
